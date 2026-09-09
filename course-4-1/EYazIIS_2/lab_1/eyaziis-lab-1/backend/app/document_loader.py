"""Извлечение текста из загружаемых файлов различных форматов."""
import io
import re

SUPPORTED_EXTENSIONS = {".txt", ".md", ".csv", ".log", ".html", ".htm", ".rtf", ".pdf", ".docx"}

_ENCODINGS = ("utf-8", "cp1251", "koi8-r", "latin-1")


class UnsupportedFormatError(ValueError):
    """Формат файла не поддерживается системой."""


def _decode(data: bytes) -> str:
    for encoding in _ENCODINGS:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _strip_html(raw: str) -> str:
    without_scripts = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", raw)
    without_tags = re.sub(r"(?s)<[^>]+>", " ", without_scripts)
    return re.sub(r"\s+", " ", without_tags).strip()


def _strip_rtf(raw: str) -> str:
    without_controls = re.sub(r"\\[a-z]+-?\d* ?", " ", raw)
    without_braces = without_controls.replace("{", " ").replace("}", " ")
    return re.sub(r"\s+", " ", without_braces).strip()


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages).strip()


def _extract_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(part for part in parts if part.strip())


def extract_text(filename: str, data: bytes) -> str:
    """Возвращает текстовое содержимое файла по его имени и байтам."""
    suffix = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFormatError(
            f"Формат {suffix or 'без расширения'} не поддерживается. "
            f"Допустимые форматы: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )

    if suffix == ".pdf":
        return _extract_pdf(data)
    if suffix == ".docx":
        return _extract_docx(data)

    text = _decode(data)
    if suffix in {".html", ".htm"}:
        return _strip_html(text)
    if suffix == ".rtf":
        return _strip_rtf(text)
    return text.strip()
